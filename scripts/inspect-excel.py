#!/usr/bin/env python3
"""Inspect Excel file structure"""

import openpyxl
import os

file_path = os.path.expanduser("~/Desktop/2025.1.5【含分类章节】HS-CIQ代码对照表-爆米hs查询.xlsx")

print("📖 Opening Excel file (full mode)...")
wb = openpyxl.load_workbook(file_path, data_only=True)

# Check a few key sheets
for sheet_name in ['HS-CIQ代码对照表', '1.1', '11.61', '11.62']:
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n=== {sheet_name} ===")
        print(f"  Max row: {ws.max_row}, Max col: {ws.max_column}")
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=3, values_only=True)):
            print(f"  Row {i+1}: {row}")

wb.close()
