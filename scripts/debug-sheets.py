#!/usr/bin/env python3
"""Debug sheet structures"""
import openpyxl
import warnings
warnings.filterwarnings("ignore")

file_path = "/Users/chq/Desktop/2025.1.5【含分类章节】HS-CIQ代码对照表-爆米hs查询.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

sheets_to_check = ["1.1", "11.61", "11.62", "15.72", "16.84"]

for sn in sheets_to_check:
    if sn in wb.sheetnames:
        ws = wb[sn]
        print(f"\n=== {sn} (max_row={ws.max_row}, max_col={ws.max_column}) ===")
        
        rows = list(ws.iter_rows(values_only=True))
        print(f"  Actual rows from iter_rows: {len(rows)}")
        
        for i, row in enumerate(rows[:5]):
            print(f"  Row {i}: {list(row)[:6]}")
        
        non_empty = sum(1 for r in rows if r and any(c is not None for c in r))
        print(f"  Non-empty rows: {non_empty}")

wb.close()
