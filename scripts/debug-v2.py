#!/usr/bin/env python3
"""Debug why chapter sheets return 0 records"""
import openpyxl
import warnings
warnings.filterwarnings('ignore')

CODE_KEYS = ['商品编码', 'hs编码', '编码', '税则号', '海关编码', '税号', 'hscode']
NAME_KEYS = ['检验检疫名称', '中文名称', '商品名称', '品名', '商品中文名']
EN_KEYS = ['检验检疫英文名称', '英文名称', '商品英文名', 'english']
CATEGORY_KEYS = ['中文描述', '分类', '章节', '类别', '章名']
CIQ_KEYS = ['ciq代码', '13位检验检疫', '检验检疫代码', 'ciq']

def col_match(val, aliases):
    if val is None:
        return False
    v = str(val).strip().lower()
    for alias in aliases:
        if alias.lower() in v:
            return True
    return False

def find_columns(header_row):
    cols = {}
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        v = str(cell).strip()
        if col_match(v, CODE_KEYS):
            if 'code' not in cols:
                cols['code'] = idx
        if col_match(v, NAME_KEYS):
            if 'name' not in cols:
                cols['name'] = idx
        if col_match(v, EN_KEYS):
            if 'en' not in cols:
                cols['en'] = idx
        if col_match(v, CATEGORY_KEYS):
            if 'category' not in cols:
                cols['category'] = idx
        if col_match(v, CIQ_KEYS):
            if 'ciq' not in cols:
                cols['ciq'] = idx
    return cols

file_path = "/Users/chq/Desktop/2025.1.5【含分类章节】HS-CIQ代码对照表-爆米hs查询.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

for sn in ['1.1', '11.61', '16.84', '20.94']:
    if sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        print(f"\n=== {sn}: {len(rows)} rows ===")
        
        # Check header detection
        for i in range(min(5, len(rows))):
            row = rows[i]
            cols = find_columns(row)
            if 'code' in cols:
                print(f"  Header found at row {i}: cols={cols}")
                print(f"  Header values: {list(row)}")
                # Check first data row
                if i+1 < len(rows):
                    data_row = rows[i+1]
                    code_idx = cols.get('code')
                    name_idx = cols.get('name')
                    print(f"  First data row: code={data_row[code_idx] if code_idx is not None else 'N/A'}, name={data_row[name_idx] if name_idx is not None else 'N/A'}")
                break
        else:
            print(f"  ❌ No header found in first 5 rows!")
            print(f"  Row 0: {rows[0] if rows else 'empty'}")
            print(f"  Row 1: {rows[1] if len(rows) > 1 else 'empty'}")

wb.close()
