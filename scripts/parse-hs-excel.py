#!/usr/bin/env python3
"""Parse commercial HS-CIQ Excel file -> JSON"""

import openpyxl
import json
import os
import warnings
warnings.filterwarnings('ignore')

file_path = os.path.expanduser("~/Desktop/2025.1.5【含分类章节】HS-CIQ代码对照表-爆米hs查询.xlsx")
output_path = os.path.expanduser("~/Desktop/hs_codes_ready.json")

print("📖 Opening Excel file...")
wb = openpyxl.load_workbook(file_path, data_only=True)

results = []
seen_codes = set()

def process_sheet(ws, sheet_name):
    count = 0
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0
    
    # Find header row
    header_idx = None
    for i, row in enumerate(rows[:5]):
        if row and row[0] and '商品编码' in str(row[0]):
            header_idx = i
            break
    
    if header_idx is None:
        return 0
    
    # Process data rows (skip header and nav rows)
    for row in rows[header_idx + 1:]:
        if not row:
            continue
        
        # Column 0: 商品编码 (HS code)
        raw_code = row[0]
        if raw_code is None:
            continue
        
        clean_code = str(raw_code).strip()
        # Remove spaces and non-digits
        clean_code = ''.join(c for c in clean_code if c.isdigit())
        
        if len(clean_code) < 4 or len(clean_code) > 13:
            continue
        
        # Column 3: 检验检疫名称 (Full Chinese product name)
        ciq_name = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ''
        
        # Column 5: 中文描述 (Category/Chapter info)
        category = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ''
        
        # Column 4: 英文名称
        desc_en = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ''
        
        # Column 2: 13位检验检疫编码
        ciq_code = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ''
        
        # Column 1: CIQ代码
        ciq_num = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ''
        
        # Skip if no meaningful name
        if not ciq_name or ciq_name == 'None':
            continue
        
        # Skip nav rows (跳转等)
        if any(k in ciq_name for k in ['跳转', '上一章节', '下一章节', '友情提示', '更多企业']):
            continue
        
        # Deduplicate
        if clean_code in seen_codes:
            continue
        seen_codes.add(clean_code)
        
        # Determine level
        level = len(clean_code)
        if level <= 4:
            level = 4
        elif level <= 6:
            level = 6
        elif level <= 8:
            level = 8
        elif level <= 10:
            level = 10
        else:
            level = 13
        
        # Build category from chapter
        if len(clean_code) >= 2:
            chap_num = int(clean_code[:2])
            if not category:
                category = f"第{chap_num}章"
        
        # Build notes
        notes = ''
        if ciq_code and len(ciq_code) > len(clean_code):
            notes += f"CIQ13: {ciq_code}"
        if ciq_num and ciq_num != 'None':
            notes += f" | CIQ: {ciq_num}" if notes else f"CIQ: {ciq_num}"
        
        results.append({
            "code": clean_code,
            "level": level,
            "description": ciq_name,
            "descriptionEn": desc_en if desc_en else ciq_name,
            "category": category if category else f"第{int(clean_code[:2])}章",
            "taxRate": None,
            "notes": notes,
        })
        count += 1
    
    return count

# Process all data sheets
meta_sheets = {'归类总则', '章节', '友情提示', '更多企业服务推荐', 'Sheet1'}
data_sheets = [s for s in wb.sheetnames if s not in meta_sheets]

print(f"Processing {len(data_sheets)} data sheets...")

for sheet_name in data_sheets:
    ws = wb[sheet_name]
    cnt = process_sheet(ws, sheet_name)
    if cnt > 0:
        print(f"  {sheet_name}: {cnt} records")

print(f"\n✅ Total: {len(results)} unique records")

# Save
with open(output_path, 'w', encoding='utf-8') as f:
    json.dump(results, f, ensure_ascii=False, indent=2)

file_size = os.path.getsize(output_path)
print(f"💾 Saved to {output_path}")
print(f"   Size: {file_size / 1024 / 1024:.1f} MB")
print(f"   Records: {len(results)}")

# Self-check
check_terms = ["床垫", "相框", "鼠标", "保温杯", "挖掘机", "大衣", "连衣裙", "咖啡", "手机"]
print(f"\n🔍 Self-check:")
for term in check_terms:
    found = [r for r in results if term in r.get('description', '')]
    if found:
        print(f"  ✅ '{term}': {len(found)} records")
        for fi in found[:2]:
            print(f"     {fi['code']} | {fi['description'][:60]}")
    else:
        print(f"  ❌ '{term}': NOT FOUND")

wb.close()
