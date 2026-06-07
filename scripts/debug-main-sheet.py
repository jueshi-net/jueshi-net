#!/usr/bin/env python3
"""Check main sheet row count"""
import openpyxl
import warnings
warnings.filterwarnings('ignore')

file_path = "/Users/chq/Desktop/2025.1.5【含分类章节】HS-CIQ代码对照表-爆米hs查询.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

ws = wb["HS-CIQ代码对照表"]
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")

rows = list(ws.iter_rows(values_only=True))
print(f"iter_rows returned: {len(rows)} rows")

# Count valid code rows
count = 0
seen = set()
for i, row in enumerate(rows):
    if not row:
        continue
    val = row[0]
    if val is not None:
        s = str(val).strip()
        digits = "".join(c for c in s if c.isdigit())
        if 4 <= len(digits) <= 13:
            count += 1
            seen.add(digits)

print(f"Rows with valid codes: {count}")
print(f"Unique codes: {len(seen)}")

wb.close()
